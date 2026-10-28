#!/usr/bin/env python3
"""YouTube Shorts 트렌드 제품 영상 수집기.

키워드별 최근 인기 쇼츠를 모아 '아웃라이어 점수'(조회수/구독자)와
'속도'(시간당 조회수), 댓글 속 구매 의도 신호를 계산해 엑셀로 저장합니다.

사용법:
    python trend_finder.py                 # 기본값(최근 7일, 미국)
    python trend_finder.py --days 3 --region GB --top 50
"""
import argparse
import json
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

API = "https://www.googleapis.com/youtube/v3"
BASE_DIR = Path(__file__).resolve().parent

# 댓글에 이런 표현이 많을수록 "사고 싶어하는" 제품
BUY_PATTERNS = [
    r"where (can|do|did) (i|you) (buy|get|find)", r"\blink\b", r"name of (the|this) product",
    r"what('?s| is) (it|this|the product) called", r"need (this|one|it)", r"want (this|one)",
    r"just (bought|ordered)", r"add(ed)? to (my )?cart", r"take my money", r"\bbuying\b",
    r"어디서 (사|구매)", r"링크", r"제품명", r"사고 싶",
]
BUY_RE = re.compile("|".join(BUY_PATTERNS), re.IGNORECASE)


def load_api_key():
    key = os.environ.get("YOUTUBE_API_KEY")
    env_file = BASE_DIR / ".env"
    if not key and env_file.exists():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("YOUTUBE_API_KEY="):
                key = line.split("=", 1)[1].strip().strip('"\'')
    if not key or "여기에" in key:
        sys.exit("YOUTUBE_API_KEY가 없습니다. README.md의 'API 키 발급' 단계를 따라 .env 파일에 넣어주세요.")
    return key


def api_get(endpoint, key, **params):
    resp = requests.get(f"{API}/{endpoint}", params={**params, "key": key}, timeout=30)
    if resp.status_code != 200:
        try:
            reason = resp.json()["error"]["errors"][0]["reason"]
        except (ValueError, KeyError, IndexError):
            reason = resp.text[:200]
        raise RuntimeError(f"{endpoint} 실패 ({resp.status_code}): {reason}")
    return resp.json()


def chunks(items, size=50):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def parse_duration(iso):
    """ISO 8601 길이(PT1M5S) → 초"""
    m = re.fullmatch(r"P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?", iso or "")
    if not m:
        return 0
    d, h, mi, s = (int(x or 0) for x in m.groups())
    return d * 86400 + h * 3600 + mi * 60 + s


def search_shorts(key, keyword, published_after, region, lang, max_results):
    data = api_get(
        "search", key, part="id", q=keyword, type="video", videoDuration="short",
        order="viewCount", publishedAfter=published_after, regionCode=region,
        relevanceLanguage=lang, maxResults=max_results,
    )
    return [item["id"]["videoId"] for item in data.get("items", [])]


def fetch_videos(key, video_ids):
    videos = {}
    for batch in chunks(video_ids):
        data = api_get("videos", key, part="snippet,statistics,contentDetails", id=",".join(batch))
        videos.update({v["id"]: v for v in data.get("items", [])})
    return videos


def fetch_channels(key, channel_ids):
    channels = {}
    for batch in chunks(channel_ids):
        data = api_get("channels", key, part="snippet,statistics", id=",".join(batch))
        channels.update({c["id"]: c for c in data.get("items", [])})
    return channels


def fetch_buy_signals(key, video_id, max_comments):
    """인기 댓글 중 구매 의도 표현이 들어간 댓글 수와 예시"""
    try:
        data = api_get("commentThreads", key, part="snippet", videoId=video_id,
                       order="relevance", maxResults=max_comments, textFormat="plainText")
    except RuntimeError:
        return 0, ""  # 댓글 비활성화 등
    hits = []
    for item in data.get("items", []):
        text = item["snippet"]["topLevelComment"]["snippet"]["textDisplay"]
        if BUY_RE.search(text):
            hits.append(text.replace("\n", " ")[:120])
    return len(hits), " | ".join(hits[:3])


def build_rows(videos, channels, keyword_map, now, max_seconds):
    rows = []
    for vid, v in videos.items():
        seconds = parse_duration(v["contentDetails"].get("duration"))
        if seconds == 0 or seconds > max_seconds:
            continue
        stats = v.get("statistics", {})
        views = int(stats.get("viewCount", 0))
        ch = channels.get(v["snippet"]["channelId"], {})
        ch_stats = ch.get("statistics", {})
        subs = 0 if ch_stats.get("hiddenSubscriberCount") else int(ch_stats.get("subscriberCount", 0))
        published = datetime.fromisoformat(v["snippet"]["publishedAt"].replace("Z", "+00:00"))
        hours = max((now - published).total_seconds() / 3600, 1)
        rows.append({
            "video_id": vid,
            "title": v["snippet"]["title"],
            "channel": v["snippet"]["channelTitle"],
            "keywords": ", ".join(sorted(keyword_map[vid])),
            "published": published.strftime("%Y-%m-%d %H:%M"),
            "seconds": seconds,
            "views": views,
            "likes": int(stats.get("likeCount", 0)),
            "comments": int(stats.get("commentCount", 0)),
            "subs": subs,
            "outlier": round(views / max(subs, 100), 1),
            "views_per_hour": round(views / hours),
            "engagement": round((int(stats.get("likeCount", 0)) + int(stats.get("commentCount", 0))) / max(views, 1) * 100, 2),
            "buy_hits": 0,
            "buy_examples": "",
        })
    return rows


def score(rows):
    """속도·아웃라이어·구매신호를 순위 기반으로 합산 (0~100)"""
    if not rows:
        return
    n = len(rows)
    for field, weight in (("views_per_hour", 0.4), ("outlier", 0.35), ("buy_hits", 0.25)):
        for rank, r in enumerate(sorted(rows, key=lambda x: x[field])):
            r["score"] = r.get("score", 0) + weight * (rank / max(n - 1, 1)) * 100
    for r in rows:
        r["score"] = round(r["score"], 1)


COLUMNS = [
    ("score", "종합점수", 9), ("title", "제목", 50), ("url", "링크", 14), ("channel", "채널", 20),
    ("views", "조회수", 12), ("views_per_hour", "시간당 조회수", 12), ("subs", "구독자", 11),
    ("outlier", "아웃라이어(조회/구독)", 12), ("engagement", "참여율%", 9), ("buy_hits", "구매의도 댓글", 10),
    ("buy_examples", "구매의도 댓글 예시", 60), ("seconds", "길이(초)", 8), ("published", "업로드(UTC)", 16),
    ("keywords", "검색 키워드", 25),
]


def save_excel(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "트렌드 쇼츠"
    ws.append([c[1] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")
    for r in rows:
        url = f"https://www.youtube.com/shorts/{r['video_id']}"
        ws.append([url if key == "url" else r[key] for key, _, _ in COLUMNS])
        link = ws.cell(row=ws.max_row, column=3)
        link.hyperlink, link.value, link.style = url, "열기", "Hyperlink"
    for i, (key, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        if key in ("views", "views_per_hour", "subs"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


JSON_FIELDS = ["video_id", "title", "channel", "keywords", "published", "seconds", "views", "likes",
               "comments", "subs", "outlier", "views_per_hour", "engagement", "buy_hits", "buy_examples", "score"]


def save_json(rows, data_dir, now, args, quota, keep):
    """홈페이지용: data/runs/<시각>.json 저장 + data/index.json 갱신(최근 keep회분만 유지)"""
    runs_dir = Path(data_dir) / "runs"
    runs_dir.mkdir(parents=True, exist_ok=True)
    stamp = now.strftime("%Y%m%d_%H%M")
    run = {
        "time": now.strftime("%Y-%m-%dT%H:%M:%SZ"), "region": args.region, "days": args.days,
        "quota": quota, "videos": [{k: r[k] for k in JSON_FIELDS} for r in rows],
    }
    (runs_dir / f"{stamp}.json").write_text(json.dumps(run, ensure_ascii=False), encoding="utf-8")

    index_path = Path(data_dir) / "index.json"
    index = json.loads(index_path.read_text(encoding="utf-8")) if index_path.exists() else []
    index = [e for e in index if e["file"] != f"runs/{stamp}.json"]
    index.insert(0, {"file": f"runs/{stamp}.json", "time": run["time"], "region": args.region, "count": len(rows)})
    for old in index[keep:]:
        (Path(data_dir) / old["file"]).unlink(missing_ok=True)
    index_path.write_text(json.dumps(index[:keep], ensure_ascii=False, indent=1), encoding="utf-8")
    return index_path


def main():
    p = argparse.ArgumentParser(description="YouTube Shorts 트렌드 제품 영상 수집기")
    p.add_argument("--keywords", default=str(BASE_DIR / "keywords.txt"), help="키워드 파일 경로")
    p.add_argument("--days", type=int, default=7, help="최근 며칠 이내 업로드 (기본 7)")
    p.add_argument("--region", default="US", help="국가 코드 (US, GB, JP, KR ...)")
    p.add_argument("--lang", default="en", help="언어 코드 (en, ja, ko ...)")
    p.add_argument("--per-keyword", type=int, default=50, help="키워드당 수집 영상 수 (최대 50)")
    p.add_argument("--max-seconds", type=int, default=180, help="이 길이(초) 이하만 쇼츠로 간주 (기본 180)")
    p.add_argument("--comment-check", type=int, default=40, help="댓글 분석할 상위 영상 수 (영상당 1유닛)")
    p.add_argument("--top", type=int, default=100, help="엑셀에 저장할 영상 수")
    p.add_argument("--out", default=str(BASE_DIR / "output"), help="엑셀 저장 폴더")
    p.add_argument("--json-dir", help="지정하면 엑셀 대신 홈페이지용 JSON을 이 폴더에 저장")
    p.add_argument("--keep", type=int, default=70, help="홈페이지에 보관할 최근 실행 횟수 (기본 70 = 약 1주일)")
    args = p.parse_args()

    key = load_api_key()
    keywords = [k.strip() for k in Path(args.keywords).read_text(encoding="utf-8").splitlines()
                if k.strip() and not k.strip().startswith("#")]
    now = datetime.now(timezone.utc)
    published_after = (now - timedelta(days=args.days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    quota = 0

    keyword_map = {}
    for kw in keywords:
        try:
            ids = search_shorts(key, kw, published_after, args.region, args.lang, min(args.per_keyword, 50))
        except RuntimeError as e:
            print(f"  [건너뜀] '{kw}': {e}")
            if "quotaExceeded" in str(e):
                break
            continue
        quota += 100
        print(f"  '{kw}': {len(ids)}개")
        for vid in ids:
            keyword_map.setdefault(vid, set()).add(kw)

    if not keyword_map:
        sys.exit("수집된 영상이 없습니다.")

    videos = fetch_videos(key, list(keyword_map))
    channels = fetch_channels(key, list({v["snippet"]["channelId"] for v in videos.values()}))
    quota += -(-len(keyword_map) // 50) + -(-len(channels) // 50)
    rows = build_rows(videos, channels, keyword_map, now, args.max_seconds)

    # 댓글 분석은 1차 점수 상위 영상에만 (할당량 절약)
    score(rows)
    rows.sort(key=lambda r: r["score"], reverse=True)
    for r in rows[:args.comment_check]:
        r["buy_hits"], r["buy_examples"] = fetch_buy_signals(key, r["video_id"], 50)
        quota += 1
    for r in rows:
        r.pop("score")
    score(rows)
    rows.sort(key=lambda r: r["score"], reverse=True)

    if args.json_dir:
        path = save_json(rows[:args.top], args.json_dir, now, args, quota, args.keep)
    else:
        out_dir = Path(args.out)
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"shorts_trend_{args.region}_{now.astimezone():%Y%m%d_%H%M}.xlsx"
        save_excel(rows[:args.top], path)

    print(f"\n쇼츠 {len(rows)}개 분석 → 상위 {min(args.top, len(rows))}개 저장: {path}")
    print(f"사용한 API 할당량: 약 {quota} / 10,000 유닛")
    for r in rows[:5]:
        print(f"  [{r['score']:5.1f}] {r['views']:>10,}회 | x{r['outlier']:<6} | {r['title'][:50]}")


if __name__ == "__main__":
    main()
